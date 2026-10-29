#!/usr/bin/env python3
"""
Turns the "Expanded Pricing Chart" spreadsheet into lib/price-chart.json, which lib/pricing.js prices quotes from.

    pip install openpyxl
    python scripts/import-pricing.py "C:/Users/User/Downloads/Expanded Pricing Chart 2026 Oct 29.xlsx"

The sheet is laid out as one block per coin size ('1" Molds $140', '2.25" Pricing - MOLDS at $229', ...), and inside
each block one table per finish group ('1" Antique Brass', '1" Shiny Nickel / Shiny Silver', ...):

    <size> <finish group>   1-5   6-15  16-25  26-40  41-50          51-74  75-99  100-299  300-499  500-749  750-999  1000-1200
    Any Color Option        12    10.85 8.15   6.4    5.4   No Color  4      3.5    2.5      2.35     2.25     2.1      1.95
                                                        1 SIDE UNLIMITED COLOR  4.5  4  3  ...
                                                        2 SIDE UNLIMITED COLOR  4.75 4.25 3.25 ...

Up to 50 coins there is one price with any color included; from 51 the price depends on whether the coins get enamel
color on no side, one side or both. To the right of the first block sit the rush fees, the shipping guide, the per-coin
add-ons and the packaging upgrades; those are copied across by cell address, so if the sheet is ever rearranged, check
the ADDRESSES below. The shipping guide's quantity columns are either a range ("1-5", "41-50") or a plain number (75,
100 ... 300), each meaning "up to that many coins"; the last column ("300+") is free or so much per coin over 300.
"""
import json
import re
import sys
from pathlib import Path

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl is needed: pip install openpyxl")

if len(sys.argv) < 2:
    sys.exit(__doc__)

SRC = Path(sys.argv[1])
OUT = Path(__file__).resolve().parent.parent / "lib" / "price-chart.json"

wb = openpyxl.load_workbook(SRC, data_only=True)
ws = wb.worksheets[0]
cell = lambda a: ws[a].value

# The finish groups on the sheet, in the order they appear in every size block, and which finishes on the quote form
# each one prices. Keys match lib/prompt.js FINISHES where they exist.
GROUPS = [
    ("antique-brass", r"antique brass", ["antique-brass"]),
    ("shiny-nickel", r"shiny nickel", ["shiny-nickel", "shiny-silver", "satin-nickel"]),
    ("copper", r"copper", ["antique-copper", "shiny-copper"]),
    ("antique-silver", r"antique silver", ["antique-silver", "satin-silver"]),
    ("gold", r"gold", ["antique-gold", "shiny-gold", "satin-gold"]),
    ("black", r"black", ["black-nickel", "matte-black", "dye-black"]),
]

SMALL_TIERS = [[1, 5], [6, 15], [16, 25], [26, 40], [41, 50]]            # columns B..F, any color included
LARGE_TIERS = [[51, 74], [75, 99], [100, 299], [300, 499], [500, 749], [750, 999], [1000, 1200]]  # columns H..N


def num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return round(float(v) + 1e-9, 2)
    m = re.search(r"\d*\.\d+|\d+", str(v))  # ".50" and "0.50" both read as half a dollar
    return round(float(m.group()), 2) if m else None


def row_vals(r, cols):
    out = []
    for c in cols:
        v = num(cell(f"{c}{r}"))
        if v is None:
            raise SystemExit(f"Expected a price at {c}{r}, found {cell(f'{c}{r}')!r}")
        out.append(v)
    return out


sizes = {}
size_order = []
current = None
for r in range(1, ws.max_row + 1):
    a = cell(f"A{r}")
    if not isinstance(a, str):
        continue
    text = a.strip()
    m = re.match(r'^([\d.]+)"\s*(?:Molds|Pricing\s*-\s*MOLDS at)\s*\$(\d+)', text, re.I)
    if m:
        current = m.group(1)
        size_order.append(current)
        sizes[current] = {"mold": int(m.group(2)), "finishes": {}}
        continue
    m = re.match(r'^([\d.]+)"\s+(.+)$', text)
    if m and cell(f"B{r}") == "1-5" and current == m.group(1):
        label = re.sub(r"\s+", " ", m.group(2)).strip()
        key = next((k for k, pat, _ in GROUPS if re.search(pat, label, re.I)), None)
        if not key:
            raise SystemExit(f"Row {r}: finish group {label!r} is not one the importer knows")
        if not str(cell(f"A{r+1}")).strip().lower().startswith("any color"):
            raise SystemExit(f"Row {r+1}: expected 'Any Color Option' under {label!r}")
        if "1 SIDE" not in str(cell(f"G{r+2}")).upper() or "2 SIDE" not in str(cell(f"G{r+3}")).upper():
            raise SystemExit(f"Rows {r+2}-{r+3}: expected the 1 SIDE / 2 SIDE color rows under {label!r}")
        sizes[current]["finishes"][key] = {
            "label": label,
            "small": row_vals(r + 1, "BCDEF"),       # 1-50 coins, any color
            "none": row_vals(r + 1, "HIJKLMN"),      # 51+ coins, no color
            "one": row_vals(r + 2, "HIJKLMN"),       # 51+ coins, color on one side
            "two": row_vals(r + 3, "HIJKLMN"),       # 51+ coins, color on both sides
        }

if not sizes:
    raise SystemExit("No size blocks found; is this the right sheet?")
for s, block in sizes.items():
    missing = [k for k, _, _ in GROUPS if k not in block["finishes"]]
    if missing:
        raise SystemExit(f'Size {s}" is missing finish groups: {missing}')

# ---- the side tables: by cell address (see the docstring) ----
ADDRESSES = {
    "rush": "P3", "shipping": "T3", "addons": "P9", "packaging": "P26",
}
for what, addr in ADDRESSES.items():
    v = str(cell(addr) or "")
    if what == "rush" and "rush" not in v.lower() or what == "shipping" and "shipping" not in v.lower() \
            or what == "addons" and "add on" not in v.lower() or what == "packaging" and "packaging" not in v.lower():
        raise SystemExit(f"Expected the {what} table heading at {addr}, found {v!r}; the sheet layout has changed")

# Rush: under 50 -> flat, 51-99 -> flat, 100+ -> +20% of the coin cost
rush = {
    "label": str(cell("Q4")).strip(),
    "flat": [[1, 50, num(cell("Q5"))], [51, 99, num(cell("Q6"))]],
    "percentFrom": 100,
    "percent": num(cell("Q7")),
    "note": "Rush does not include shipping upgrades.",
}

# Shipping guide: rows by size range, columns by quantity. Each column heading is the most coins it covers: a range
# like "1-5" or "41-50" (its upper end), or a plain 75, 100 ... 300. The column after the last number ("300+") is what
# shipping costs past it. Headings are read from row 4 until the "+" column, so the sheet can add or drop columns.
ship_cols, ship_letters = [], []
for c in range(21, 40):  # U onwards
    head = cell(f"{get_column_letter(c)}4")
    if head is None or "+" in str(head):
        over_col = get_column_letter(c)
        break
    m = re.match(r"^\s*(?:\d+\s*-\s*)?(\d+)\s*$", str(head))
    if not m:
        raise SystemExit(f"Shipping guide heading {head!r} at {get_column_letter(c)}4 is not a quantity")
    ship_cols.append(int(m.group(1)))
    ship_letters.append(get_column_letter(c))
else:
    raise SystemExit("Shipping guide: no '300+' column found after the quantity columns")
if ship_cols != sorted(ship_cols) or not ship_cols:
    raise SystemExit(f"Shipping guide columns are not in rising order: {ship_cols}")
shipping = []
for r in range(5, 14):
    label = str(cell(f"T{r}") or "").strip()
    m = re.match(r'^([\d.]+)"(?:\s*to\s*([\d.]+)")?', label)
    if not m:
        break
    lo, hi = m.group(1), m.group(2) or m.group(1)
    flat = []
    for c in ship_letters:
        v = cell(f"{c}{r}")
        flat.append(0 if str(v).strip().upper() == "FREE" else num(v))
    over = cell(f"{over_col}{r}")
    over_s = str(over).strip()
    shipping.append({
        "sizes": [lo, hi], "upTo": ship_cols, "flat": flat,
        # over the last column: free, or so many dollars per coin (".50¢ per coin" and "$0.75 per coin" both mean dollars:
        # the sheet runs .50, .55, .65, $0.75, $0.85 ...)
        "overPerCoin": 0 if over_s.upper() == "FREE" else num(over_s),
    })

# Add-ons: Item / Per Coin / Setup
addons = []
r = 11
while True:
    item = cell(f"P{r}")
    if not isinstance(item, str) or not item.strip() or item.strip().lower().startswith("packaging"):
        break
    label = re.sub(r"\s+", " ", item).strip()
    key = re.sub(r"[^a-z0-9]+", "-", label.lower()).strip("-")
    addons.append({"key": key, "label": label, "perCoin": num(cell(f"Q{r}")), "setup": num(cell(f"R{r}")) or 0})
    r += 1

# Packaging: Item / 0-74 / 75-500 / + S&H each
packaging = []
r = 28
while True:
    item = cell(f"P{r}")
    if not isinstance(item, str) or not item.strip():
        break
    label = re.sub(r"\s+", " ", item).strip().replace("'", '"')
    key = re.sub(r"[^a-z0-9]+", "-", label.lower()).strip("-")
    packaging.append({"key": key, "label": label, "under75": num(cell(f"Q{r}")), "from75": num(cell(f"R{r}")), "shippingEach": num(cell(f"S{r}")) or 0})
    r += 1

def group_label(key):
    # "Antique Brass", "Shiny Nickel / Shiny Silver", ...: as the first size block names it, typos tidied
    label = sizes[size_order[0]]["finishes"][key]["label"]
    return re.sub(r"\s*/\s*", " / ", label.replace("Nickell", "Nickel")).strip()


chart = {
    "source": SRC.name,
    "title": str(cell("A1") or "").strip(),
    "sizes": size_order,
    "smallTiers": SMALL_TIERS,
    "largeTiers": LARGE_TIERS,
    "finishGroups": [{"key": k, "finishes": f, "label": group_label(k)} for k, _, f in GROUPS],
    "prices": sizes,
    "addons": addons,
    "packaging": packaging,
    "rush": rush,
    "shipping": shipping,
}
OUT.write_text(json.dumps(chart, indent=2) + "\n", encoding="utf-8")
print(f"Wrote {OUT}: {len(size_order)} sizes, {len(GROUPS)} finish groups, {len(addons)} add-ons, {len(packaging)} packaging options, {len(shipping)} shipping rows")
